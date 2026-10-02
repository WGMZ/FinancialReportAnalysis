"""Markdown -> Word converter and builder for REPORT_zh.docx / REPORT_en.docx.

Supports: headings, paragraphs, pipe tables, bullet / numbered lists (2 levels), block quotes,
display and inline TeX (rendered with sub/superscripts), bold, italic, inline code, images with captions.
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.join(os.path.dirname(__file__), "..")
OUT = os.path.join(BASE, "output")
SRC = os.path.join(OUT, "doc_src")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x55, 0x55, 0x55)
PAGE_W_CM = 17.2

GREEK = {"Delta": "Δ", "tau": "τ", "sum": "Σ", "cdot": "·", "times": "×", "rho": "ρ", "alpha": "α", "beta": "β", "sigma": "σ", "mu": "μ"}
STRIP = ("big", "Big", "bigg", "Bigg", "left", "right", "bigl", "bigr")


def set_font(run, east="Microsoft YaHei", ascii_="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = ascii_
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), ascii_)
    rf.set(qn("w:hAnsi"), ascii_)
    rf.set(qn("w:eastAsia"), east)
    rf.set(qn("w:cs"), ascii_)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


# ---------------------------------------------------------------- TeX
def _group(s, i):
    """Return (content, next_index) of a {...} group starting at s[i]=='{' or a single token."""
    if i >= len(s):
        return "", i
    if s[i] == "{":
        d, j = 1, i + 1
        while j < len(s) and d:
            d += (s[j] == "{") - (s[j] == "}")
            j += 1
        return s[i + 1:j - 1], j
    if s[i] == "\\":
        m = re.match(r"\\[A-Za-z]+", s[i:])
        if m:
            return m.group(0), i + len(m.group(0))
    return s[i], i + 1


def tex_tokens(s, sub=False, sup=False, upright=False):
    out, i = [], 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+)", s[i:])
            if m:
                name = m.group(1)
                i += len(m.group(0))
                if name in STRIP:
                    continue
                if name == "text" or name == "mathrm":
                    g, i = _group(s, i)
                    out += tex_tokens(g, sub, sup, True)
                    continue
                if name == "frac":
                    a, i = _group(s, i)
                    b, i = _group(s, i)
                    out += [("(", sub, sup, upright)] + tex_tokens(a, sub, sup, upright) + [(")/(", sub, sup, upright)] + tex_tokens(b, sub, sup, upright) + [(")", sub, sup, upright)]
                    continue
                out.append((GREEK.get(name, name), sub, sup, True if name in GREEK and name in ("sum", "cdot", "times") else upright))
                continue
            nxt = s[i + 1] if i + 1 < len(s) else ""
            i += 2
            out.append((" " if nxt == "," else nxt, sub, sup, True))
            continue
        if c in "_^":
            g, i = _group(s, i + 1)
            out += tex_tokens(g, c == "_", c == "^", upright)
            continue
        if c in "{}":
            i += 1
            continue
        out.append((c, sub, sup, upright or not c.isalpha()))
        i += 1
    return out


def add_math(par, tex, size=None, color=None, base_east="Microsoft YaHei"):
    toks = tex_tokens(tex)
    merged = []
    for t, sb, sp, up in toks:
        if merged and merged[-1][1:] == (sb, sp, up):
            merged[-1] = (merged[-1][0] + t, sb, sp, up)
        else:
            merged.append((t, sb, sp, up))
    for t, sb, sp, up in merged:
        r = par.add_run(t)
        set_font(r, east=base_east, ascii_="Cambria Math", size=size, italic=not up, color=color)
        if sb:
            r.font.subscript = True
        if sp:
            r.font.superscript = True


# ---------------------------------------------------------------- inline
MATH_RE = re.compile(r"\$([A-Za-z\\][A-Za-z0-9_^{}\\(),.=+\-% ]*?)\$")
TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\$[A-Za-z\\][A-Za-z0-9_^{}\\(),.=+\-% ]*?\$|(?<![*\w])\*[^*\s][^*]*?\*(?![*\w]))")


def _is_math(inner):
    return len(inner) == 1 or any(ch in inner for ch in "\\_^")


def add_inline(par, text, size=10.5, bold=False, italic=False, color=None):
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(par, part[2:-2], size, True, italic, color)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            set_font(r, east="Microsoft YaHei", ascii_="Consolas", size=size - 1, color=RGBColor(0x8B, 0x1A, 0x1A))
        elif part.startswith("$") and part.endswith("$") and len(part) > 2 and _is_math(part[1:-1]):
            add_math(par, part[1:-1], size=size, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(par, part[1:-1], size, bold, True, color)
        else:
            r = par.add_run(part)
            set_font(r, size=size, bold=bold or None, italic=italic or None, color=color)


# ---------------------------------------------------------------- document
def new_doc(lang):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.18
    for name, sz, before, after in (("Title", 22, 0, 8), ("Heading 1", 17, 14, 8), ("Heading 2", 14, 14, 6), ("Heading 3", 12, 10, 4), ("Heading 4", 11, 8, 3)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = NAVY
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn("w:" + a), "Microsoft YaHei" if a == "eastAsia" else "Calibri")
        for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            if rf.get(qn("w:" + a)) is not None:
                del rf.attrib[qn("w:" + a)]
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    footer_page_numbers(doc, lang)
    return doc


def footer_page_numbers(doc, lang):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 " if lang == "zh" else "Page ")
    set_font(r, size=9, color=GREY)
    for kind in ("PAGE",):
        run = p.add_run()
        set_font(run, size=9, color=GREY)
        for t, txt in (("begin", None), (None, kind), ("end", None)):
            if t:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), t)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)
    if lang == "zh":
        r = p.add_run(" 页")
        set_font(r, size=9, color=GREY)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("美股金融板块 2026 Q3 财报季前瞻" if lang == "zh" else "US Financials - Q3 2026 earnings preview")
    set_font(r, size=8.5, color=GREY)


def para_border_left(par, color="1F4E79", fill="F2F5F9"):
    ppr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    b.append(left)
    ppr.append(b)
    shade(ppr, fill)


def add_table(doc, rows):
    header = rows[0]
    body = rows[1:]
    n = len(header)
    t = doc.add_table(rows=1 + len(body), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    fs = 9 if n <= 4 else (8.5 if n <= 6 else (8 if n <= 8 else 7.5))
    plain = lambda s: re.sub(r"[*`$]", "", s)  # noqa: E731
    lens = [max([len(plain(r[c])) if c < len(r) else 0 for r in rows] + [3]) for c in range(n)]
    w = [min(max(l, 4), 46) ** 0.8 for l in lens]
    tot = sum(w)
    widths = [PAGE_W_CM * x / tot for x in w]
    for ci, wd in enumerate(widths):
        t.columns[ci].width = Cm(wd)
    for ri, row in enumerate([header] + body):
        tr = t.rows[ri]
        if ri == 0:
            trpr = tr._tr.get_or_add_trPr()
            h = OxmlElement("w:tblHeader")
            h.set(qn("w:val"), "true")
            trpr.append(h)
        cant = OxmlElement("w:cantSplit")
        tr._tr.get_or_add_trPr().append(cant)
        for ci in range(n):
            cell = tr.cells[ci]
            cell.width = Cm(widths[ci])
            txt = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            short = len(plain(txt)) <= 16
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ci > 0 and short) or ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, txt, size=fs, bold=(ri == 0), color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
            tcpr = cell._tc.get_or_add_tcPr()
            if ri == 0:
                shade(tcpr, "1F4E79")
            elif ri % 2 == 0:
                shade(tcpr, "F2F5F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines):
    rows = []
    for ln in lines:
        s = ln.strip().strip("|")
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", s)]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def add_image(doc, path, caption):
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(PAGE_W_CM - 0.4))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    add_inline(c, caption, size=9, italic=True, color=GREY)


def render(doc, md, base_dir):
    lines = md.split("\n")
    i = 0
    first_h1 = True
    while i < len(lines):
        ln = lines[i]
        s = ln.rstrip()
        if not s.strip() or re.fullmatch(r"-{3,}", s.strip()):
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl, text = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and first_h1:
                first_h1 = False
                p = doc.add_paragraph(style="Title")
                add_inline(p, text, size=22, bold=True, color=NAVY)
            else:
                if lvl == 1:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                p = doc.add_paragraph(style=f"Heading {lvl}")
                add_inline(p, text, size={1: 17, 2: 14, 3: 12, 4: 11}[lvl], bold=True, color=NAVY)
            i += 1
            continue
        if s.startswith("|"):
            blk = []
            while i < len(lines) and lines[i].startswith("|"):
                blk.append(lines[i])
                i += 1
            add_table(doc, parse_table(blk))
            continue
        if s.startswith(">"):
            paras, cur = [], []
            while i < len(lines) and lines[i].startswith(">"):
                t = lines[i][1:].strip()
                if t:
                    cur.append(t)
                elif cur:
                    paras.append(" ".join(cur))
                    cur = []
                i += 1
            if cur:
                paras.append(" ".join(cur))
            for t in paras:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                para_border_left(p)
                add_inline(p, t, size=9.5, color=RGBColor(0x33, 0x33, 0x33))
            continue
        if s.startswith("$$"):
            tex = s.strip("$ ")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            add_math(p, tex, size=12)
            i += 1
            continue
        im = re.match(r"^!\[(.*)\]\((.*)\)$", s)
        if im:
            add_image(doc, os.path.normpath(os.path.join(base_dir, im.group(2))), im.group(1))
            i += 1
            continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if lm:
            indent = len(lm.group(1))
            lvl = 0 if indent == 0 else (1 if indent < 5 else 2)
            marker, text = lm.group(2), lm.group(3)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.65 + 0.7 * lvl)
            pf.first_line_indent = Cm(-0.5)
            pf.space_after = Pt(3)
            bullet = ("•", "–", "·")[lvl] if marker in "-*" else marker
            r = p.add_run(bullet + "\t")
            set_font(r, size=10.5, color=NAVY if marker in "-*" else None, bold=marker not in "-*")
            pf.tab_stops.add_tab_stop(Cm(0.65 + 0.7 * lvl))
            add_inline(p, text)
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\||>|\$\$|!\[|\s*([-*]|\d+\.)\s)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))


def build(lang):
    read = lambda p: open(p, encoding="utf-8").read()  # noqa: E731
    front = read(os.path.join(SRC, f"front_{lang}.md"))
    supp = read(os.path.join(SRC, f"supp_{lang}.md"))
    body = read(os.path.join(OUT, "REPORT.md" if lang == "zh" else "REPORT_en.md"))
    body = re.sub(r"^# .*\n", "", body, count=1)
    doc = new_doc(lang)
    render(doc, front + "\n\n" + body + "\n\n" + supp, SRC)
    cp = doc.core_properties
    cp.title = "美股金融板块 2026 年 Q3 财报季前瞻(图文版)" if lang == "zh" else "US Financials Ahead of the Q3 2026 Earnings Season (Illustrated Edition)"
    cp.author = "Research"
    out = os.path.join(OUT, f"REPORT_{lang}.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    for lang in (sys.argv[1:] or ["zh", "en"]):
        print(build(lang))
